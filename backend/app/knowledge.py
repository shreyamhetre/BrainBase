"""The Knowledge library: turn an uploaded PDF or Word file, or a GitLab file link, into a document.

Parsing produces Markdown so the normal chunker keeps headings: Word headings become #/##/###,
PDF pages become "## Page n". Storage goes through the shared store (ingestion/store.py), so
added documents are chunked, embedded, permission-tagged and searchable like everything else.
"""

import io
import re
import urllib.parse
import uuid
from dataclasses import dataclass

import psycopg

from app.auth import User
from app.connectors.clients import ConnectorError, client
from app.ingestion.store import SourceDoc, sync_documents
from app.vault import decrypt

MAX_BYTES = 10 * 1024 * 1024          # 10 MB per file
MAX_TEXT_CHARS = 2_000_000            # about 500k tokens: refuse anything bigger
PDF = "application/pdf"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
TEXT_SUFFIXES = (".md", ".markdown", ".txt")


class KnowledgeError(Exception):
    """A message the person can act on, e.g. "This PDF has no text (is it a scan?)"."""


# ---------- who may add, and with whom they may share ----------

def can_add(conn: psycopg.Connection, user: User) -> bool:
    if user.is_admin:
        return True
    return bool(conn.execute("SELECT 1 FROM roles WHERE name = ANY(%s) AND can_add_documents LIMIT 1",
                             (user.roles,)).fetchone())


def shareable_roles(user: User) -> list[str]:
    """You can share with "everyone" or with roles you hold yourself, never with roles you don't have."""
    return ["all", *user.roles]


# ---------- parsing ----------

def pdf_to_markdown(data: bytes) -> str:
    from pypdf import PdfReader
    try:
        reader = PdfReader(io.BytesIO(data))
        pages = [(n, (page.extract_text() or "").strip()) for n, page in enumerate(reader.pages, 1)]
    except Exception as e:  # pypdf raises many kinds of errors for broken files
        raise KnowledgeError(f"Couldn't read this PDF: {e}") from None
    text = "\n\n".join(f"## Page {n}\n\n{t}" for n, t in pages if t)
    if not text:
        raise KnowledgeError("This PDF has no selectable text (is it a scan?). Upload a text PDF or a Word file.")
    return text


def docx_to_markdown(data: bytes) -> str:
    from docx import Document
    try:
        doc = Document(io.BytesIO(data))
    except Exception as e:
        raise KnowledgeError(f"Couldn't read this Word file: {e}") from None
    lines = []
    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        style = (p.style.name if p.style is not None else "").lower()
        level = re.search(r"heading (\d)", style)
        if style == "title":
            lines.append(f"# {text}")
        elif level:
            lines.append(f"{'#' * min(int(level.group(1)), 3)} {text}")
        elif "list" in style:
            lines.append(f"- {text}")
        else:
            lines.append(text)
    for table in doc.tables:  # keep tables as Markdown rows
        rows = [[c.text.strip().replace("|", "/") for c in row.cells] for row in table.rows]
        if rows:
            lines.append("| " + " | ".join(rows[0]) + " |")
            lines.append("|" + " --- |" * len(rows[0]))
            lines += ["| " + " | ".join(r) + " |" for r in rows[1:]]
    text = "\n\n".join(lines)
    if not text:
        raise KnowledgeError("This Word file is empty.")
    return text


def parse(file_name: str, content_type: str, data: bytes) -> tuple[str, str]:
    """(markdown, content type) for an uploaded or fetched file."""
    if len(data) > MAX_BYTES:
        raise KnowledgeError("Files can be up to 10 MB.")
    name = file_name.lower()
    if name.endswith(".pdf") or content_type == PDF:
        text, kind = pdf_to_markdown(data), PDF
    elif name.endswith(".docx") or content_type == DOCX:
        text, kind = docx_to_markdown(data), DOCX
    elif name.endswith(TEXT_SUFFIXES):
        text, kind = data.decode("utf-8", errors="replace"), "text/markdown"
    elif name.endswith(".doc"):
        raise KnowledgeError("Old .doc files aren't supported: save it as .docx and upload again.")
    else:
        raise KnowledgeError("Upload a PDF or Word (.docx) file, or link a Markdown, text, PDF or Word file in GitLab.")
    if len(text) > MAX_TEXT_CHARS:
        raise KnowledgeError("This document is too long to index (over 2 million characters).")
    return text, kind


def title_from(file_name: str, text: str) -> str:
    heading = re.search(r"^# (.+)$", text, re.M)
    if heading:
        return heading.group(1).strip()[:200]
    stem = re.sub(r"\.[a-z0-9]+$", "", file_name, flags=re.I)
    return re.sub(r"[_-]+", " ", stem).strip()[:200] or "Untitled document"


# ---------- storing ----------

def store(conn: psycopg.Connection, *, source: str, external_id: str, title: str, url: str, text: str,
          allowed_roles: list[str], user: User, metadata: dict, access_id: str | None = None,
          section: str = "Knowledge library") -> int:
    """Chunk, embed and save one document; returns its id."""
    doc = SourceDoc(source=source, external_id=external_id, title=title, url=url, body=text,
                    allowed_roles=allowed_roles, section=section, doc_type=metadata.get("kind", "document"),
                    owners=[user.name], metadata=metadata, access_id=access_id)
    with conn.transaction():
        stats = sync_documents(conn, [doc], "source = %s AND external_id = %s", (source, external_id), full=False)
        if stats.changed == 0 and stats.unchanged == 0:
            raise KnowledgeError("No text could be indexed from this document.")
        doc_id = conn.execute("UPDATE documents SET added_by = %s, updated_at = coalesce(updated_at, now()) "
                              "WHERE source = %s AND external_id = %s RETURNING id",
                              (user.id, source, external_id)).fetchone()[0]
    return doc_id


def new_upload_id() -> str:
    return str(uuid.uuid4())


# ---------- GitLab links ----------

@dataclass
class GitLabFile:
    project_path: str      # e.g. engineering
    ref: str
    file_path: str


LINK = re.compile(r"^https?://[^/]+/(?P<group>[^/]+(?:/[^/]+)*?)/(?P<project>[^/]+)/-/(?:blob|raw)/(?P<ref>[^/]+)/(?P<path>.+)$")


def parse_gitlab_link(url: str, group: str) -> GitLabFile:
    match = LINK.match(url.strip().split("?")[0].split("#")[0])
    if not match:
        raise KnowledgeError("Paste a link to a file in GitLab, like https://gitlab.com/<group>/<project>/-/blob/main/docs/policy.md")
    if match.group("group").lower() != group.lower():
        raise KnowledgeError(f"Only files in the connected GitLab group ({group}) can be added.")
    return GitLabFile(match.group("project"), match.group("ref"), urllib.parse.unquote(match.group("path")))


def fetch_gitlab_file(conn: psycopg.Connection, link: GitLabFile) -> tuple[bytes, dict]:
    """The file's bytes and the project it lives in, through the GitLab connector's token."""
    row = conn.execute("SELECT id::text, token_encrypted, config FROM connectors WHERE kind = 'gitlab' AND enabled "
                       "ORDER BY created_at LIMIT 1").fetchone()
    if not row:
        raise KnowledgeError("Connect GitLab on the Connectors page first.")
    connector_id, token, config = row
    project = conn.execute("SELECT id::text, container_id, container_name, allowed_roles FROM source_access "
                           "WHERE connector_id = %s AND container_name = %s", (connector_id, link.project_path)).fetchone()
    if not project:
        raise KnowledgeError(f"The project '{link.project_path}' isn't connected. Check the Connectors page.")
    api = client("gitlab", decrypt(token), config)
    path = urllib.parse.quote(link.file_path, safe="")
    try:
        data, _ = _get_raw(api, f"/projects/{project[1]}/repository/files/{path}/raw?ref={urllib.parse.quote(link.ref, safe='')}")
    except ConnectorError as e:
        raise KnowledgeError(f"GitLab couldn't return that file: {e}") from None
    return data, {"access_id": project[0], "project": project[2], "allowed_roles": list(project[3])}


def _get_raw(api, path: str) -> tuple[bytes, dict]:
    import urllib.error
    import urllib.request
    req = urllib.request.Request(api.base + path, headers=api.headers)
    try:
        with urllib.request.urlopen(req, timeout=30) as r:
            return r.read(), dict(r.headers)
    except urllib.error.HTTPError as e:
        raise ConnectorError(f"HTTP {e.code}") from None
